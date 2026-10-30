import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/arthurklein/Documents/Work/2026Nov/VR HS')
BUILD = ROOT / '.character-dossier-build'
OUT = ROOT / 'production/project-files/VR_High_School_Character_Register_v1_1.docx'
faculty = json.loads((BUILD / 'faculty.json').read_text())
motion = json.loads((BUILD / 'motion.json').read_text())
voices = json.loads((BUILD / 'voices_v1_1.json').read_text())
for f, voice in zip(faculty, voices):
    f['prompt'] = voice['prompt']
    f['pace'] = voice['pace']
faculty[1]['test'][3][1] = 'A rune machine doubles its input, then adds three. Can you write a rule for any input?'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.60)
sec.left_margin = sec.right_margin = Inches(.65)
sec.header_distance = sec.footer_distance = Inches(.24)
styles = doc.styles
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Arial')
    style.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Arial')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.line_spacing = 1.04
styles['Normal'].paragraph_format.space_after = Pt(4)
styles['Title'].font.size = Pt(29)
styles['Title'].font.bold = True
styles['Title'].paragraph_format.space_after = Pt(12)
styles['Heading 1'].font.size = Pt(20)
styles['Heading 1'].font.bold = True
styles['Heading 1'].paragraph_format.space_after = Pt(7)
styles['Heading 2'].font.size = Pt(12)
styles['Heading 2'].font.bold = True
styles['Heading 2'].paragraph_format.space_before = Pt(9)
styles['Heading 2'].paragraph_format.space_after = Pt(4)
styles['Heading 3'].font.size = Pt(10.5)
styles['Heading 3'].font.bold = True
for name in ['Heading 1', 'Heading 2', 'Heading 3']:
    styles[name].paragraph_format.keep_with_next = True

header = sec.header.paragraphs[0]
header.text = 'VR HIGH SCHOOL   CHARACTER REGISTER'
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Character register v1.1   /   ').font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def p(text='', boldlead=None, size=None, after=None):
    para = doc.add_paragraph()
    if boldlead:
        para.add_run(boldlead).bold = True
        para.add_run(text)
    else:
        para.add_run(text)
    if size:
        for run in para.runs: run.font.size = Pt(size)
    if after is not None: para.paragraph_format.space_after = Pt(after)
    return para

pending_page = False
def h(text, level=2):
    global pending_page
    para = doc.add_paragraph(text, 'Heading '+str(level))
    if pending_page:
        para.paragraph_format.page_break_before = True
        pending_page = False
    return para
def page():
    global pending_page
    pending_page = True

def table(headers, rows, widths, font=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for i, title in enumerate(headers): t.rows[0].cells[i].text = title
    for values in rows:
        cells = t.add_row().cells
        for i, text in enumerate(values): cells[i].text = str(text)
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'1D303E' if ri==0 else ('F1F5F6' if ri%2 else 'FFFFFF'));tcpr.append(sh)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1)
                para.paragraph_format.line_spacing=1.03
                if ci==0: para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size=Pt(font);run.font.bold=(ri==0)
                    run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
        trpr=row._tr.get_or_add_trPr()
        split=OxmlElement('w:cantSplit');trpr.append(split)
        if ri==0: trpr.append(OxmlElement('w:tblHeader'))
    return t

doc.add_paragraph('VR High School Character Register', 'Title')
p('Voice design animation prompts and test dialogue', size=15, after=11)
p('Version 1.1   /   29 October 2026   /   Angular Arcade character direction', size=10, after=12)
p('This working reference keeps the 13 faculty identities, 14 proposed voices and animation prompts together. Use the stable character IDs to name assets and record approved choices. The ettin has two separately designed voices, F05A and F05B.')
h('ElevenLabs voice workflow')
p('Open Voices, then My Voices, Add a new voice and Voice Design. Paste the Voice Design prompt from the relevant dossier. Put the separate voice preview text into Text to preview. Generate candidates, audition them and save the selected voice. [1]')
p('Use the saved voice in Text to Speech for the test dialogue. Generate only that character’s spoken lines. Role labels and student lines are rehearsal directions, not text for the teacher voice. For the ettin, keep Head A and Head B separate and alternate them in the edit.')
p('Record the saved voice ID, speech model and the actual settings used. Pace ranges are acting targets. Use the same short scene to compare revisions. Avoid assuming that a numeric setting or expressive tag works across every speech model.')
h('Animation workflow')
p('Use ElevenLabs Image & Video in Video mode with a reference-capable model available in your account. Supply the matching individual character reference. Combine the Shared animation brief, the character’s Identity lock and one clip prompt. Generate one action per clip. [2]')
p('Choose the closest supported duration to the target in this register, then trim as needed. Generate silent movement tests first. Add the approved teacher audio to the speaking clip. Use an audio or lip-sync option only when the selected model supports it. Check creature mouths and appendages manually. [2]')
h('Production record')
p('All voice and motion assets begin as Planned. No voice IDs, generated takes or approvals have been assigned in this register. Update the status to In review, Approved or Revise after inspecting a take. Keep superseded takes under their existing version numbers.')
p('Video generations are motion references. They do not supply a game-ready mesh, skeleton or animation rig. Use the same clip directions when an animator builds the real-time character.')

page()
h('Faculty index and selection tracker',1)
p('Identity follows the teacher roster. Performance ages, delivery and animation behavior are proposed creative choices. F05A and F05B belong to the same ettin body. “TBD” means a selection has not yet been made.', after=9)
rows=[]
for i,f in enumerate(faculty):
    ident=f'F{i+1:02d}'
    voice='F05A + F05B' if i==4 else ident
    rows.append([ident,f['subject'].title(),f['creature'].title(),f['gender'],voice+'\nID TBD','Planned'])
table(['ID','Subject','Character','Gender','Voice record','Assets'],rows,[.43,1.12,1.60,.77,1.27,1.01],font=9.2)
h('Naming convention')
p('Use character_action_v### for motion and character_line_v### for dialogue. Examples: F01_idle_v001.mp4, F01_welcome_v001.wav, F05A_welcome_v001.wav and F05B_question_v001.wav. Keep the source prompt with each take.')
h('Continuity decisions')
p('History remains female, following the teacher table where older classroom prose differs. The headmaster is female. The mimic is non-binary. English emphasizes sentence craft, while Language Arts emphasizes audience and storytelling. No personal names have been assigned.')

page()
h('Voice casting map',1)
p('Use species, physical scale and demeanor together when auditioning. These are stylized acting choices for fantasy creatures. Size shapes vocal weight and resonance, while anatomy shapes texture and articulation. Each teacher still speaks clearly enough for a student to follow a lesson.')
cast_rows = []
for i, voice in enumerate(voices):
    if i == 4:
        cast_rows.append(['F05A','Enormous ettin head','Earthy gravelly bass-baritone','Grounded and deliberate'])
        cast_rows.append(['F05B','Same giant body','Ringing higher baritone with body','Curious and quicker'])
    else:
        cast_rows.append([f'F{i+1:02d}',voice['physical'],voice['anchor'],voice['demeanor']])
table(['Voice','Body and species','Vocal signature','Demeanor'],cast_rows,[.55,2.05,2.75,1.85],font=9.5)
h('Shared audition line')
p('Take your time. First, tell me what you notice. That is a useful beginning. Shall we try another way?')
p('Generate this same text with every candidate after the character preview. Compare vocal weight, texture and rhythm at similar playback volume. The yeti should feel dense and soft, the dragon ringing and expansive, the lich light and dry. The spider should sound fine and careful, while the goblin sounds compact and nimble.')
p('Then use the character’s own test dialogue to check reassurance and instruction. Start with moderate creature coloration. If a texture obscures words, reduce that one trait while preserving the character’s register and phrasing.')

page()
h('Shared animation brief',1)
p('Copy this brief first, append the character’s Identity lock, then append exactly one of that character’s clip prompts. Use the same reference file and prompt revision for each comparison take.')
h('Shared prompt')
p('Animate only the single teacher shown in the supplied character reference. Preserve the Angular Arcade appearance: faceted cuboid forms, matte materials, the same face, anatomy, clothing, accessories and colors. Keep every limb and appendage consistent. Use a locked camera at a comfortable eye level with the full silhouette in frame. Keep the classroom background and lighting stable. Perform one small, readable action with gentle acceleration and a settled finish. The teacher feels patient, welcoming and interested in the learner. Keep the clip silent for the first movement test. No captions, logos, camera shake, cuts, sudden approach, additional characters or redesign.')
h('Clip library')
table(['Action','Target','Result'],[
    ['Idle','6 seconds','Small neutral loop with matching start and end poses'],
    ['Listen','6 seconds','Acknowledge once and hold a quiet listening pose'],
    ['Greet','6 seconds','Welcome gesture followed by a return to rest'],
    ['Teach','8 seconds','One clear demonstration with a visible pause'],
    ['Travel','6 seconds','Short movement across frame with a settled stop'],
    ['Speak','One audio line','One teacher line with restrained speech motion']
],[.82,.95,5.43],font=10)
p('Durations describe intended edited clips. Model duration options vary. Speech clips follow the actual audio length. Split long lines into separate clips when needed. A loop still requires frame-by-frame review and may need an edit to match its endpoints.',after=6).paragraph_format.space_before=Pt(6)
h('Review before approving')
p('Compare the take with the reference. Check limb count, costume, props and face stability. Look for sliding feet, clipping, distorted hands and changing body volume. Confirm that the main gesture is readable and that the teacher settles while listening. Keep the camera and lighting consistent across takes.')
p('For reduced motion, use a still or very small idle pose, quiet appendages and no displacement effect. The history travel prompt uses a solid silhouette by default. Characters remain outside the learner’s personal space.')
h('Speech synchronization')
p('Human-style lip sync may distort a beholder, chuul, dragon, slime, mimic or lich. Use a restrained jaw, lid, mouthpart or head gesture when necessary. For the ettin, only the head assigned to the current audio line moves its mouth.')

for i,(f,m) in enumerate(zip(faculty,motion)):
    page()
    ident=f'F{i+1:02d}'
    h(ident+' '+f['subject'].title()+' '+f['creature'].title(),1)
    p(f["gender"]+' / '+f['className']+' / '+f['location']+' / '+f['pace'],size=10,after=4)
    p('Voice IDs: F05A TBD / F05B TBD    TTS model and settings: TBD    Status: Planned' if i==4 else 'Voice ID: TBD    TTS model and settings: TBD    Voice status: Planned',size=9.2,after=2)
    p('Motion model and reference: TBD    Latest take: None    Motion status: Planned',size=9.2,after=5)
    h('Voice Design prompt')
    prompt=f['prompt']
    if i==11:
        prompt=prompt.replace(' Choose the voice for character fit rather than a gender stereotype.','')
    p(prompt,after=4)
    h('Voice preview text')
    if i==4:
        p('F05A: Welcome, explorers. The map shows two paths to the tower. One crosses the ridge, and one follows the river. Let us compare the distance and the slope. Show me a feature on the map that supports your route.',size=10,after=3)
        p('F05B: One route is shorter. Does that make it easier? I would like to check those contour lines before we choose. We can disagree and still study the same evidence. What would change your mind about the route?',size=10,after=3)
    else:
        preview=' '.join(line for role,line in f['test'] if role!='STUDENT')
        if i==0:
            preview=preview.replace('Then it is doing its job.', 'A first draft does not need to be perfect.')
        p(preview,size=10,after=4)
    h('Test dialogue')
    for role,line in f['test']:
        label=role.replace(' · ',' ').title()+': '
        p(line,boldlead=label,size=10,after=2)
    h('Animation prompts')
    p(m['lock'],boldlead='Identity lock: ',size=10,after=4)
    for action in ['idle','listen','greet','teach','travel','speak']:
        dur='match one audio line' if action=='speak' else ('8 seconds' if action=='teach' else '6 seconds')
        p(m[action],boldlead=f'{ident}_{action} ({dur}): ',size=10,after=3)
    p('Review notes: __________________________________________________________________',size=9,after=0)

page()
h('Generation and approval log',1)
p('Use one row per take. Add rows as production continues. Keep voice candidates and motion clips under the same character ID. Attach the exact prompt and settings to the exported asset or its companion record.')
table(['Date','Character and take','Model or voice ID','Status and reviewer','Notes'],[
    ['','','','',''],['','','','',''],['','','','',''],['','','','','']
],[.64,1.55,1.50,1.45,1.56],font=9.5)
h('Prompt revision log')
table(['Rev','Date','Change'],[['1.0','29 Oct 2026','Initial register of 13 faculty, 14 voices, 78 motion prompts and rehearsal scenes.'],['1.1','29 Oct 2026','More distinct voice prompts shaped by creature anatomy, physical scale and demeanor. Added a casting map and common audition line.']], [.65,1.05,5.50],font=10)
h('Sources and working files')
p('Teacher identities: docs/teachers.md. Visual direction: docs/art/concept-direction.md and assets/concepts/2026-10-29/02-faculty-angular-arcade.png. These paths are relative to the VR HS project folder.')
p('Companion presentation: VR High School Faculty Dossiers, the 21-slide deck created in this chat. This register carries forward its cast, voice direction and test scenes, including an introductory algebra warm-up for Math.')
p('ElevenLabs documentation checked 29 October 2026. Model availability and controls vary by account and selected model. Use the current interface when recording settings.')
p('[1] ElevenLabs Voice Design\nhttps://elevenlabs.io/docs/eleven-creative/voices/voice-design',size=9.5)
p('[2] ElevenLabs Image and Video\nhttps://elevenlabs.io/docs/eleven-creative/playground/image-video',size=9.5)

doc.core_properties.title='VR High School Character Register'
doc.core_properties.subject='Character tracking voice design animation prompts and test dialogue'
doc.core_properties.author='VR High School'
doc.core_properties.keywords='VRHS, character register, ElevenLabs, voice design, animation'
for el in doc.element.iter():
    for border in list(el.findall(qn('w:pBdr'))):
        el.remove(border)
for style in styles:
    for el in style.element.iter():
        for border in list(el.findall(qn('w:pBdr'))):
            el.remove(border)
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
print('Characters:',len(faculty),'Voices:',14,'Motion prompts:',len(motion)*6)
